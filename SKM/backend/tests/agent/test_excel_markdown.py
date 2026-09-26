"""同一 XLSX 原値に対する削除線・fill・座標と未評価書式を検証する。"""

from __future__ import annotations

from copy import copy
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Color, Font, GradientFill, PatternFill
from openpyxl.styles.fills import Stop

from skillmind.agent.binary_text import convert_excel_to_markdown, render_excel_markdown
from skillmind.agent.excel_markdown import _CellTokens, _Colors, _compact_ranges, markdown_literal


def _bytes(book: Workbook) -> bytes:
    """外部接続のない合成 workbook を原 byte 列として確定する。"""
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def _book() -> Workbook:
    """値と座標が明確な一枚のテスト表を作る。"""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Cases"
    sheet.append(["ID", "Action", "Expected"])
    sheet.append(["TC-1", "old action", "old result"])
    sheet.append(["TC-2", "new action", "new result"])
    return book


def _convert(book: Workbook) -> str:
    """本番の MarkItDown 経路まで通す。"""
    return render_excel_markdown(".xlsx", _bytes(book))


@pytest.mark.parametrize(
    "rgb", ["FFFF0000", "FFFFFF00", "FF00FF00", "FFD9D9D9", "FF1234AB", "FFFFFFFF"]
)
def test_any_solid_fill_preserves_raw_and_resolved_color(rgb: str) -> None:
    """灰色以外と明示白色も保持し、色を理由に行を除外しない。"""
    book = _book()
    book["Cases"]["B2"].fill = PatternFill("solid", fgColor=rgb)
    markdown = _convert(book)
    assert "| B2 | F1 | unspecified |" in markdown
    assert f"rgb={rgb}" in markdown and "#" + rgb[-6:] in markdown
    assert "```json" not in markdown
    assert "old action" in markdown and "new action" in markdown
    assert "~~old action~~" not in markdown


def test_whole_cell_strike_and_literal_tildes_are_distinct() -> None:
    """本物の削除線のみ描画し、値中の ~~ を書式へ昇格させない。"""
    book = _book()
    book["Cases"]["B2"].font = Font(strike=True)
    book["Cases"]["B3"] = "~~literal~~"
    markdown = _convert(book)
    assert "~~old action~~" in markdown
    assert r"\~\~literal\~\~" in markdown
    assert "| B2 | F1 | true |" in markdown


def test_rich_text_keeps_exact_struck_span_and_explicit_false() -> None:
    """セル font の削除線を run の false が解除し、日本語/Unicode 座標を保持する。"""
    book = _book()
    cell = book["Cases"]["B2"]
    cell.font = Font(strike=True)
    cell.value = CellRichText(
        TextBlock(InlineFont(strike=False), "入力 "),
        TextBlock(InlineFont(strike=True), "旧😀"),
        TextBlock(InlineFont(strike=False), " 新しい値"),
    )
    markdown = _convert(book)
    assert "~~旧😀~~" in markdown
    assert "~~入力" not in markdown and "~~新しい" not in markdown
    assert "| B2 | " + markdown_literal("[3,5)") + " |" in markdown


def test_empty_and_duplicate_headers_keep_original_row_and_columns() -> None:
    """pandas header 推論の Unnamed/.1 を作らず、空行の原行番号も保持する。"""
    book = _book()
    sheet = book["Cases"]
    sheet["A1"] = None
    sheet["B1"] = "same"
    sheet["C1"] = "same"
    sheet.insert_rows(2)
    markdown = _convert(book)
    assert "| Excel row | A | B | C |" in markdown
    assert "| 1 |  | same | same |" in markdown
    assert "| 2 |  |  |  |" in markdown
    assert "| 3 | TC-1 | old action | old result |" in markdown
    assert "Unnamed" not in markdown and "same.1" not in markdown


def test_compaction_preserves_holes_and_blank_color_legends() -> None:
    """同書式の連続矩形だけを縮約し、空白の凡例や離れた cell も記録する。"""
    book = _book()
    sheet = book["Cases"]
    for coordinate in ("A2", "B2", "A3", "B3", "D8"):
        sheet[coordinate].fill = PatternFill("solid", fgColor="FF00AABB")
    markdown = _convert(book)
    assert "| A2:B3, D8 | F1 | unspecified |" in markdown
    assert _compact_ranges([(1, 1), (1, 3), (2, 1), (2, 2), (2, 3)]) == ["A1", "C1", "A2:C2"]


@pytest.mark.parametrize("tint,expected", [(0, "#FF0000"), (0.5, "#FF8080"), (-0.5, "#800000")])
def test_theme_color_uses_workbook_theme_and_hls_tint(tint: float, expected: str) -> None:
    """固定の Office palette を仮定せず、原 theme と tint を残す。"""
    book = _book()
    book.loaded_theme = b"""<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:themeElements><a:clrScheme name="test"><a:lt1><a:srgbClr val="112233"/></a:lt1><a:dk1><a:srgbClr val="334455"/></a:dk1><a:accent1><a:srgbClr val="FF0000"/></a:accent1></a:clrScheme></a:themeElements></a:theme>"""
    book["Cases"]["B2"].fill = PatternFill("solid", fgColor=Color(theme=4, tint=tint))
    markdown = _convert(book)
    assert "theme=4" in markdown and expected in markdown
    if tint:
        assert f"tint={tint}" in markdown


def test_custom_indexed_palette_and_reserved_system_colors() -> None:
    """workbook 独自 palette を使い、system 色を default に置換しない。"""
    book = _book()
    book._colors = ["FF000000", "FFABCDEF"]
    sheet = book["Cases"]
    sheet["A2"].fill = PatternFill("solid", fgColor=Color(indexed=1))
    sheet["B2"].fill = PatternFill("solid", fgColor=Color(indexed=64))
    sheet["C2"].fill = PatternFill("solid", fgColor=Color(auto=True))
    markdown = _convert(book)
    assert markdown_literal("indexed=1; #ABCDEF; RESOLVED") in markdown
    assert "indexed=64; UNRESOLVED" in markdown
    assert "auto=True; UNRESOLVED" in markdown


def test_missing_theme_keeps_unresolved_number() -> None:
    """theme が欠けても値を捨てず、色解決だけが未完了と分かる。"""
    book = _book()
    colors = _Colors(book)
    value = colors.describe(Color(theme=4, tint=0.25))
    assert value == {"type": "theme", "value": 4, "tint": 0.25, "resolution": "UNRESOLVED"}


def test_pattern_and_gradient_are_not_flattened_to_one_color() -> None:
    """fg/bg と gradient stop を区別し、StyleProxy 経由でも欠落させない。"""
    book = _book()
    sheet = book["Cases"]
    sheet["B2"].fill = PatternFill("darkGrid", fgColor="FFFF0000", bgColor="FF0000FF")
    sheet["B3"].fill = GradientFill(
        degree=45, stop=[Stop(Color(rgb="FFFF0000"), 0), Stop(Color(rgb="FF00FF00"), 1)]
    )
    markdown = _convert(book)
    assert "darkGrid" in markdown and "background: rgb=FF0000FF" in markdown
    assert '"type":"gradient"' in markdown and '"degree":45' in markdown
    assert '"position":0' in markdown and '"position":1' in markdown


def test_conditional_fill_and_strike_are_explicitly_unevaluated() -> None:
    """条件が偽/真でも求値せず、範囲・式・差分書式を元 data として残す。"""
    book = _book()
    sheet = book["Cases"]
    sheet["B2"].fill = PatternFill("solid", fgColor="FFFFFF00")
    sheet.conditional_formatting.add(
        "A2:C3",
        CellIsRule(
            operator="equal",
            formula=['"obsolete"'],
            fill=PatternFill("solid", fgColor="FFFF0000"),
            font=Font(strike=True),
            stopIfTrue=True,
        ),
    )
    sheet.conditional_formatting.add(
        "A5:A8",
        ColorScaleRule(start_type="min", start_color="FF0000", end_type="max", end_color="00FF00"),
    )
    markdown = _convert(book)
    assert "NOT_EVALUATED" in markdown
    assert "| A2:C3 |" in markdown and "| A5:A8 |" in markdown
    assert "obsolete" in markdown and '"strike":true' in markdown
    assert "~~old action~~" not in markdown
    assert "#FFFF00" in markdown


def test_merged_anchors_and_dimension_defaults_are_not_expanded() -> None:
    """結合範囲と default style は記録するが、複写や全行走査を行わない。"""
    book = _book()
    sheet = book["Cases"]
    sheet.merge_cells("B2:C2")
    sheet["B2"].fill = PatternFill("solid", fgColor="FFFF0000")
    sheet.row_dimensions[4].fill = PatternFill("solid", fgColor="FF00FF00")
    sheet.column_dimensions["D"].font = Font(strike=True)
    markdown = _convert(book)
    assert "Merged: B2:C2" in markdown and "| B2 | F1 | unspecified |" in markdown
    assert "| row | 4 |" in markdown and "| column | D" in markdown


def test_original_special_values_are_not_html_or_markdown_instructions() -> None:
    """改行・CR・HTML・pipe・literal token の内容を壊さず一回だけ展開する。"""
    book = _book()
    sheet = book["Cases"]
    sheet["B2"] = "  001|<script>\r\nNA\t~~old~~  "
    sheet["B3"] = "SKMCELLTOKEN0END"
    # stdlib XML writer は CR を生で出力し、再読時に LF へ正規化される。
    # fixture 自体が原 CR を保存するよう文字参照にし、converter の断言は弱めない。
    raw = _bytes(book)
    target = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"\r", b"&#13;")
            archive.writestr(item.filename, data)
    prepared = target.getvalue()
    verified = load_workbook(BytesIO(prepared), data_only=True, rich_text=True)
    try:
        assert verified["Cases"]["B2"].value == "  001|<script>\r\nNA\t~~old~~  "
    finally:
        verified.close()
    markdown = render_excel_markdown(".xlsx", prepared)
    assert "&#32;&#32;001\\|&lt;script&gt;&#13;<br>NA&#9;\\~\\~old\\~\\~&#32;&#32;" in markdown
    assert "<script>" not in markdown
    assert "SKMCELLTOKEN0END" in markdown


def test_tokens_preserve_empty_and_whitespace_only_struck_runs() -> None:
    """空文字は delimiter を増やさず、空白だけの値は残す。"""
    tokens = _CellTokens()
    assert tokens.segment("", True) == ""
    assert tokens.restore(tokens.segment(" \t ", True)) == "&#32;&#9;&#32;"


def test_style_reader_failure_is_visible_in_saved_markdown(monkeypatch: pytest.MonkeyPatch) -> None:
    """追加書式読取だけの不具合で元値を捨てず、無書式の成功とも主張しない。"""
    from skillmind.agent import excel_markdown

    def fail(*args: Any, **kwargs: Any) -> None:
        """値読取とは独立な追加 reader の失敗。"""
        raise ValueError("unsupported formatting")

    monkeypatch.setattr(excel_markdown, "load_workbook", fail)
    markdown = _convert(_book())
    assert "old action" in markdown
    assert "NOT_INSPECTED" in markdown
    assert "STYLE_EXTRACTION_UNAVAILABLE" in markdown


def test_shared_string_rich_text_is_not_limited_to_inline_strings() -> None:
    """sharedStrings の run も同じ削除線で表現する。"""
    raw = _bytes(_book())
    target = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(target, "w", ZIP_DEFLATED) as result:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(
                    b'<c r="B2" t="inlineStr"><is><t>old action</t></is></c>',
                    b'<c r="B2" t="s"><v>0</v></c>',
                )
            elif item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
                )
            result.writestr(item.filename, data)
        result.writestr(
            "xl/sharedStrings.xml",
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><r><rPr><strike/></rPr><t>old</t></r><r><t> action</t></r></si></sst>',
        )
    markdown = render_excel_markdown(".xlsx", target.getvalue())
    assert "~~old~~" in markdown and " action" in markdown


async def test_style_conversion_uses_real_bounded_subprocess() -> None:
    """親子 process 境界を越えて書式が残り、原 byte 列は変わらない。"""
    book = _book()
    book["Cases"]["B2"].font = Font(strike=True)
    book["Cases"]["B2"].fill = PatternFill("solid", fgColor="FFFF0000")
    data = _bytes(book)
    before = copy(data)
    result = await convert_excel_to_markdown("sample.xlsx", data)
    assert data == before and result.profile == "excel-styles/v2"
    assert "~~old action~~" in result.markdown and "#FF0000" in result.markdown


async def test_published_artifact_contains_styles_and_keeps_source_identity() -> None:
    """Agent が再転記しなくても、保存 Artifact と Evidence に同じ変換が残る。"""
    from skillmind.agent.document_provider import DocumentConvertProvider
    from tests.agent.test_document_conversion import _conversion_context
    from tests.agent.test_document_provider import _FakeSource
    from tests.documents.fakes import document_content

    book = _book()
    book["Cases"]["B2"].fill = PatternFill("solid", fgColor="FFFF0000")
    content = document_content(_bytes(book), name="styles.xlsx")
    context = _conversion_context(content)
    source = _FakeSource(project_id=context.project_id, content=content)
    result = await DocumentConvertProvider(source).execute(
        context, {"path": "specs/styles.xlsx", "publish_artifact": True}
    )
    artifact = result.evidence[-1].artifact
    assert artifact is not None and artifact.content.decode("utf-8") == result.response["markdown"]
    assert result.response["document"]["checksum"] == content.checksum
    assert result.evidence[0].metadata["conversion_profile"] == "excel-styles/v2"
    assert result.response["markdown_checksum"] == artifact.checksum
    assert "#FF0000" in artifact.content.decode("utf-8")


async def test_legacy_xls_values_remain_available_without_claiming_style_support() -> None:
    """旧形式の既存変換は継続するが、色や取消線を調べたとは表現しない。"""
    from tests.agent.test_document_conversion import ROOT

    raw = (ROOT / "backend/tests/fixtures/documents/conversion-sample.xls").read_bytes()
    result = await convert_excel_to_markdown("legacy.xls", raw)
    assert result.profile == "excel-values/v1"
    assert "NOT inspected" in result.markdown and "正常終了" in result.markdown


def test_repeated_cell_style_is_resolved_once(monkeypatch: pytest.MonkeyPatch) -> None:
    """同色の長い表でセル数に比例した fill 複製/解色を繰り返さない。"""
    book = _book()
    for row in range(2, 502):
        book["Cases"].cell(row, 2, "item").fill = PatternFill("solid", fgColor="FFABCDEF")
    calls = 0
    original = _Colors.fill

    def counted(self: _Colors, fill: Any) -> dict[str, Any]:
        """解色の実呼出しを計測し、処理内容は変えない。"""
        nonlocal calls
        calls += 1
        return original(self, fill)

    monkeypatch.setattr(_Colors, "fill", counted)
    markdown = _convert(book)
    assert calls == 2
    assert "| B2:B501 | F1 | unspecified |" in markdown


def test_plain_rich_text_does_not_expand_style_report_but_strike_override_remains() -> None:
    """無装飾 rich text の空配列は省略し、基底削除線の解除を明示する。"""
    book = _book()
    for row in range(2, 102):
        book["Cases"].cell(row, 2).value = CellRichText(TextBlock(InlineFont(), "text"))
    cell = book["Cases"]["C2"]
    cell.font = Font(strike=True)
    cell.value = CellRichText(TextBlock(InlineFont(strike=False), "active"))
    markdown = _convert(book)
    facts = markdown.split("### Excel style facts", 1)[1]
    assert "| C2 | none |" in facts
    assert "| B2 | none |" not in facts
    assert "~~active~~" not in markdown
    assert "```json" not in facts and len(facts) < 1200


def test_many_disjoint_ranges_are_readable_without_giant_json_lines() -> None:
    """穴を潰さず範囲行を分け、色定義は一度だけ載せる。"""
    book = _book()
    for row in range(2, 302, 2):
        book["Cases"].cell(row, 2, "item").fill = PatternFill("solid", fgColor="FFD9D9D9")
    markdown = _convert(book)
    facts = markdown.split("### Excel style facts", 1)[1]
    assert facts.count("rgb=FFD9D9D9") == 1
    assert "B2:B300" not in facts
    assert max(map(len, facts.splitlines())) < 300
    assert "B300" in facts
