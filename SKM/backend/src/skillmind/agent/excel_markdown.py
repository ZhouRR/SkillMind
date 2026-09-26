"""Excel の静的書式を、業務判断せず Markdown と原座標へ投影する。

値の読取は既存 pandas 経路が担当する。本モジュールは原 workbook を保存・再計算せず、
削除線の文字範囲と任意の fill を保持する。条件付き書式・table style は表示結果ではない。
"""

from __future__ import annotations

import colorsys
import json
import re
import warnings
from collections import defaultdict
from collections.abc import Callable, Iterable, Mapping, Sequence
from copy import copy
from dataclasses import dataclass, field
from html import escape
from io import BytesIO
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.styles.colors import Color
from openpyxl.styles.fills import Fill, GradientFill
from openpyxl.utils.cell import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from skillmind.core.hashing import canonical_json

STYLE_PROFILE = "excel-styles/v2"
VALUE_PROFILE = "excel-values/v1"
_THEME_KEYS = (
    "lt1",
    "dk1",
    "lt2",
    "dk2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
)
_DRAWING_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_MARKER = re.compile(r"SKMCELLTOKEN[0-9]+END")


def markdown_literal(text: str) -> str:
    """元値を Markdown 命令や HTML と解釈させず、改行/CR/端の空白も区別する。"""
    text = escape(text, quote=False)
    text = re.sub(r"([\\`*_{}\[\]()#+!|~])", r"\\\1", text)
    text = text.replace("\r", "&#13;").replace("\n", "<br>").replace("\t", "&#9;")
    text = re.sub(r"^ +| +$", lambda match: "&#32;" * len(match[0]), text)
    return text


@dataclass
class _CellTokens:
    """HTML 変換による値の strip/空白圧縮を避ける一回限りの置換表。"""

    values: list[str] = field(default_factory=list)

    def add(self, text: str) -> str:
        """値自体は HTML converter に渡さず、元順序で識別する。"""
        token = f"SKMCELLTOKEN{len(self.values)}END"
        self.values.append(markdown_literal(text))
        return token

    def segment(self, text: str, strike: bool) -> str:
        """削除線の外に端の空白を置き、GFM の delimiter 条件を満たす。"""
        if not text:
            return ""
        if not strike or not text.strip():
            return self.add(text)
        left = len(text) - len(text.lstrip())
        right = len(text.rstrip())
        return (
            (self.add(text[:left]) if left else "")
            + "<del>"
            + self.add(text[left:right])
            + "</del>"
            + (self.add(text[right:]) if right < len(text) else "")
        )

    def restore(self, markdown: str) -> str:
        """一回の regex 置換なので、元文字列中の同名 token を再解釈しない。"""
        return _MARKER.sub(lambda match: self.values[int(match[0][12:-3])], markdown)


class _Colors:
    """workbook 固有の theme/palette と原色表現を保持する。"""

    def __init__(self, book: Workbook) -> None:
        """system color は lastClr の観測値のみを使い、OS 色を推測しない。"""
        self.palette = tuple(book._colors)
        self.theme: dict[int, str] = {}
        if book.loaded_theme:
            try:
                root = ElementTree.fromstring(book.loaded_theme)
                scheme = root.find(f".//{_DRAWING_NS}clrScheme")
                if scheme is not None:
                    for index, key in enumerate(_THEME_KEYS):
                        node = scheme.find(f"{_DRAWING_NS}{key}")
                        if node is not None and len(node):
                            color = node[0]
                            # 色変換を持つ theme 定義は、base color として偽装しない。
                            if not len(color):
                                value = (
                                    color.get("val")
                                    if color.tag.endswith("}srgbClr")
                                    else (
                                        color.get("lastClr")
                                        if color.tag.endswith("}sysClr")
                                        else None
                                    )
                                )
                                if value is not None:
                                    self.theme[index] = value
            except ElementTree.ParseError:
                # セル値は取得済み。欠損 theme は原番号と未解決として報告する。
                pass

    def describe(self, color: Color) -> dict[str, Any]:
        """解決できない色を白色/無色へ変換せず、tint 前後と原値を残す。"""
        result: dict[str, Any] = {"type": color.type, "value": color.value, "tint": color.tint}
        value: str | None = None
        if color.type == "rgb":
            value = str(color.rgb)
        elif color.type == "theme":
            value = self.theme.get(int(color.theme))
        elif color.type == "indexed":
            index = int(color.indexed)
            if 0 <= index < min(64, len(self.palette)):
                value = str(self.palette[index])
        if value is None or re.fullmatch(r"(?:[0-9a-fA-F]{6}|[0-9a-fA-F]{8})", value) is None:
            result["resolution"] = "UNRESOLVED"
            return result
        base = value[-6:].upper()
        result["base_rgb"] = "#" + base
        if color.tint:
            r, g, b = (int(base[offset : offset + 2], 16) / 255 for offset in (0, 2, 4))
            h, luminance, saturation = colorsys.rgb_to_hls(r, g, b)
            luminance = (
                luminance * (1 + color.tint)
                if color.tint < 0
                else luminance * (1 - color.tint) + color.tint
            )
            channels = colorsys.hls_to_rgb(h, luminance, saturation)
            result["rgb"] = "#" + "".join(f"{round(channel * 255):02X}" for channel in channels)
            result["resolution"] = "DERIVED_HLS"
        else:
            result["rgb"] = "#" + base
            result["resolution"] = "RESOLVED"
        return result

    def fill(self, fill: Fill) -> dict[str, Any]:
        """solid と pattern/gradient を区別し、単一背景色へ丸めない。"""
        fill = copy(fill)
        if isinstance(fill, GradientFill):
            return {
                "type": "gradient",
                "gradient": fill.type,
                "degree": fill.degree,
                "left": fill.left,
                "right": fill.right,
                "top": fill.top,
                "bottom": fill.bottom,
                "stops": [
                    {"position": stop.position, "color": self.describe(stop.color)}
                    for stop in fill.stop
                ],
            }
        # openpyxl の StyleProxy も同じ公開 fill 属性を提供する。
        pattern = getattr(fill, "patternType", None)
        if pattern is None:
            return {"type": "none"}
        return {
            "type": "pattern",
            "pattern": pattern,
            "foreground": self.describe(getattr(fill, "fgColor")),
            **(
                {"background": self.describe(getattr(fill, "bgColor"))}
                if pattern != "solid"
                else {}
            ),
        }


def _compact_ranges(cells: Iterable[tuple[int, int]]) -> list[str]:
    """実際に存在する同書式セルだけを水平/垂直結合し、穴を含む矩形を作らない。"""
    rows: dict[int, list[int]] = defaultdict(list)
    for row, column in cells:
        rows[row].append(column)
    rectangles: list[list[int]] = []
    previous: dict[tuple[int, int], list[int]] = {}
    for row in sorted(rows):
        columns = sorted(set(rows[row]))
        intervals: list[tuple[int, int]] = []
        start = end = columns[0]
        for column in columns[1:]:
            if column == end + 1:
                end = column
            else:
                intervals.append((start, end))
                start = end = column
        intervals.append((start, end))
        current: dict[tuple[int, int], list[int]] = {}
        for interval in intervals:
            rectangle = previous.get(interval)
            if rectangle is not None and rectangle[2] == row - 1:
                rectangle[2] = row
            else:
                rectangle = [row, interval[0], row, interval[1]]
                rectangles.append(rectangle)
            current[interval] = rectangle
        previous = current
    return [
        f"{get_column_letter(c1)}{r1}"
        + (f":{get_column_letter(c2)}{r2}" if (r1, c1) != (r2, c2) else "")
        for r1, c1, r2, c2 in rectangles
    ]


def _style_facts(sheet: Worksheet, colors: _Colors) -> dict[str, Any]:
    """疎な実セルを走査し、空白の色凡例や広い row/column 設定も展開せず残す。"""
    groups: dict[str, list[tuple[int, int]]] = defaultdict(list)
    rich_strikes: list[dict[str, Any]] = []
    styles: dict[int, dict[str, Any]] = {}
    for (row, column), cell in sorted(sheet._cells.items()):
        # merge の非 anchor は合成セルであり、原セルの書式とは主張しない。
        if cell.__class__.__name__ == "MergedCell":
            continue
        # 同じ workbook の同じ immutable style は一度だけ解色する。
        detail = styles.get(cell.style_id)
        if detail is None:
            detail = {"fill": colors.fill(cell.fill)}
            if cell.font.strike is not None:
                detail["base_font_strike"] = cell.font.strike
            styles[cell.style_id] = detail
        if detail["fill"] != {"type": "none"} or cell.font.strike is not None:
            groups[canonical_json(detail)].append((row, column))
        if isinstance(cell.value, CellRichText):
            spans = []
            offset = 0
            for text, strike in _rich_segments(cell.value, bool(cell.font.strike)):
                if strike and text:
                    spans.append([offset, offset + len(text)])
                offset += len(text)
            # 削除線なしの rich text を列挙しない。base=true の明示解除だけは空範囲を残す。
            if spans or cell.font.strike:
                rich_strikes.append({"cell": cell.coordinate, "strike_spans": spans})
    facts: dict[str, Any] = {
        "profile": STYLE_PROFILE,
        "cell_styles": [
            {"ranges": _compact_ranges(coords), **json.loads(detail)}
            for detail, coords in sorted(groups.items())
        ],
    }
    if rich_strikes:
        facts["rich_text"] = rich_strikes
        facts["rich_text_offsets"] = (
            "Zero-based Unicode code points, end exclusive; overrides base_font_strike."
        )
    defaults = []
    for kind, dimensions in (("row", sheet.row_dimensions), ("column", sheet.column_dimensions)):
        for key, dimension in sorted(dimensions.items(), key=lambda entry: str(entry[0])):
            if dimension.has_style:
                defaults.append(
                    {
                        "axis": kind,
                        "key": str(key),
                        **(
                            {"min": dimension.min, "max": dimension.max} if kind == "column" else {}
                        ),
                        "fill": colors.fill(dimension.fill),
                        "strike": dimension.font.strike,
                    }
                )
    if defaults:
        facts["dimension_defaults"] = defaults
        facts["dimension_note"] = "Row/column defaults are not expanded; explicit cell styles win."
    merges = sorted(str(item) for item in sheet.merged_cells.ranges)
    if merges:
        facts["merged_ranges"] = merges
        facts["merge_note"] = (
            "Value and static style refer to the anchor; covered cells are not copied."
        )
    conditional = []
    for region in sheet.conditional_formatting:
        rules = []
        for rule in sheet.conditional_formatting[region]:
            entry: dict[str, Any] = {
                "type": rule.type,
                "priority": rule.priority,
                "stop_if_true": rule.stopIfTrue,
                "operator": rule.operator,
                "formulas": list(rule.formula or ()),
            }
            if rule.dxf is not None:
                if rule.dxf.fill is not None:
                    entry["fill"] = colors.fill(rule.dxf.fill)
                if rule.dxf.font is not None:
                    entry["strike"] = rule.dxf.font.strike
            if rule.colorScale is not None:
                entry["color_scale"] = [colors.describe(color) for color in rule.colorScale.color]
            rules.append(entry)
        conditional.append(
            {"range": str(region.sqref), "evaluation": "NOT_EVALUATED", "rules": rules}
        )
    if conditional:
        facts["conditional_formats"] = conditional
    if sheet.tables:
        facts["table_styles"] = [
            {"range": table.ref, "evaluation": "NOT_EVALUATED"} for table in sheet.tables.values()
        ]
    if sheet.sheet_state != "visible":
        facts["sheet_state"] = sheet.sheet_state
    return facts


def _rich_segments(value: Any, strike: bool) -> list[tuple[str, bool]]:
    """削除線の明示 false を尊重し、同じ書式の隣接 run のみまとめる。"""
    pieces: list[tuple[str, bool]] = []
    for part in value if isinstance(value, CellRichText) else (str(value),):
        text = part.text if isinstance(part, TextBlock) else str(part)
        flag = (
            part.font.strike
            if isinstance(part, TextBlock) and part.font.strike is not None
            else strike
        )
        if pieces and pieces[-1][1] == flag:
            pieces[-1] = (pieces[-1][0] + text, flag)
        else:
            pieces.append((text, flag))
    return pieces


def _color_text(color: Mapping[str, Any]) -> str:
    """原色指定と解決結果を短く併記し、未解決と tint の意味を保持する。"""
    parts = [f"{color['type']}={color['value']}"]
    if color.get("tint"):
        parts.append(f"tint={color['tint']}")
    if "rgb" in color:
        parts.append(str(color["rgb"]))
    if color.get("base_rgb") != color.get("rgb") and "base_rgb" in color:
        parts.append(f"base={color['base_rgb']}")
    parts.append(str(color["resolution"]))
    return "; ".join(parts)


def _fill_text(fill: Mapping[str, Any]) -> str:
    """通常の塗りを文章化し、稀な gradient の詳細も失わない。"""
    if fill["type"] == "none":
        return "none"
    if fill["type"] == "pattern":
        text = f"{fill['pattern']}: {_color_text(fill['foreground'])}"
        if "background" in fill:
            text += f" / background: {_color_text(fill['background'])}"
        return text
    return json.dumps(fill, ensure_ascii=False, separators=(",", ":"))


def _facts_table(headers: Sequence[str], rows: Iterable[Sequence[str]]) -> str:
    """原文字を escape して、外部内容が Markdown の列や HTML を作らない表を返す。"""
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines.extend("| " + " | ".join(markdown_literal(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def _style_markdown(facts: Mapping[str, Any]) -> str:
    """書式事実を範囲表へ投影する。共通色は一回、巨大な range 配列は短い行へ分ける。"""
    lines = ["### Excel style facts", "", f"Profile: {STYLE_PROFILE}"]
    if "style_status" in facts:
        lines.append(str(facts["style_status"]))
    styles = facts.get("cell_styles", [])
    fills: dict[str, str] = {}
    for entry in styles:
        key = canonical_json(entry["fill"])
        fills.setdefault(key, f"F{len(fills) + 1}")
    if fills:
        lines.extend(["", _facts_table(
            ["Fill", "Definition"],
            ((label, _fill_text(json.loads(key))) for key, label in fills.items()),
        )])
    rows = []
    for entry in styles:
        ranges = entry["ranges"]
        strike = entry.get("base_font_strike")
        for start in range(0, len(ranges), 12):
            rows.append((
                ", ".join(ranges[start:start + 12]), fills[canonical_json(entry["fill"])],
                "true" if strike is True else "false" if strike is False else "unspecified",
            ))
    if rows:
        lines.extend(["", _facts_table(["Cells", "Fill", "Base strike"], rows)])
    else:
        lines.append("No explicit cell fill/strike entries." if "style_status" not in facts
                     else "Cell formatting was not inspected.")
    if "rich_text" in facts:
        lines.extend(["", str(facts["rich_text_offsets"]), _facts_table(
            ["Rich text cell", "Strike spans"],
            ((item["cell"], ", ".join(f"[{a},{b})" for a, b in item["strike_spans"]) or "none")
             for item in facts["rich_text"]),
        )])
    if "dimension_defaults" in facts:
        lines.extend(["", str(facts["dimension_note"]), _facts_table(
            ["Axis", "Key / bounds", "Fill", "Strike"],
            ((item["axis"],
              item["key"] + (f" ({item['min']}:{item['max']})" if "min" in item else ""),
              _fill_text(item["fill"]), json.dumps(item["strike"]))
             for item in facts["dimension_defaults"]),
        )])
    if "merged_ranges" in facts:
        lines.extend(["", str(facts["merge_note"])])
        merges = facts["merged_ranges"]
        lines.extend("Merged: " + ", ".join(merges[i:i + 12]) for i in range(0, len(merges), 12))
    # 条件式など複雑な例外だけ JSON 表記を残す。内容を評価済みとして扱わない。
    for key in ("conditional_formats", "table_styles"):
        if key in facts:
            lines.extend(["", key + " (NOT_EVALUATED)", _facts_table(
                ["Range", "Details"],
                ((item["range"], json.dumps({k: v for k, v in item.items() if k != "range"},
                  ensure_ascii=False, separators=(",", ":"))) for item in facts[key]),
            )])
    if "sheet_state" in facts:
        lines.append("Sheet state: " + str(facts["sheet_state"]))
    return "\n".join(lines)


def _sheet_table(
    rows: Sequence[Sequence[str]],
    sheet: Worksheet | None,
    html_to_markdown: Callable[[str], str],
) -> str:
    """原行/列を表中に示し、先頭行を header 推論で改名/消失させない。"""
    tokens = _CellTokens()
    width = max((len(row) for row in rows), default=0)
    output = [
        "<table><thead><tr><th>Excel row</th>"
        + "".join(f"<th>{get_column_letter(column)}</th>" for column in range(1, width + 1))
        + "</tr></thead><tbody>"
    ]
    for row_number, values in enumerate(rows, start=1):
        output.append(f"<tr><td>{row_number}</td>")
        for column, value in enumerate(values, start=1):
            cell = sheet._cells.get((row_number, column)) if sheet is not None else None
            strike = bool(cell.font.strike) if cell is not None else False
            segments = [(value, strike)]
            if cell is not None and isinstance(cell.value, CellRichText):
                # rich text と値読取の対応は完全一致。違えば別の文字へ線を移さない。
                if str(cell.value) != value:
                    raise ValueError("Rich text does not match the source cell value")
                segments = _rich_segments(cell.value, strike)
            output.append(
                "<td>" + "".join(tokens.segment(text, flag) for text, flag in segments) + "</td>"
            )
        output.append("</tr>")
    output.append("</tbody></table>")
    return tokens.restore(html_to_markdown("".join(output)))


def render_styled_sheets(
    data: bytes,
    sheets: Mapping[str, Any],
    html_to_markdown: Callable[[str], str],
) -> str:
    """同じ XLSX bytes から書式を取得し、一文書に本文と事実を格納する。

    既存の値読取が成功した workbook で追加の書式読取だけが失敗した場合は、明示的な
    未読注記を添える。内容を失敗扱いにしたり、書式無しと偽装したりしない。
    """
    book: Workbook | None = None
    limitations = [
        "Conditional formatting and table styles are not evaluated; static fill is not "
        "the final displayed color. The converter infers no exclusions; interpret style facts using the source legend and business rules."
    ]
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            book = load_workbook(BytesIO(data), data_only=True, keep_links=False, rich_text=True)
        if caught:
            limitations.append(
                "Workbook reader reported unsupported features; visual fidelity is incomplete."
            )
    except (ValueError, TypeError, KeyError, IndexError, OSError, ElementTree.ParseError):
        limitations.append(
            "STYLE_EXTRACTION_UNAVAILABLE: only cell values were read; do not infer "
            "absence of strikethrough or background formatting."
        )
    try:
        colors = _Colors(book) if book is not None else None
        sections = []
        for name, frame in sheets.items():
            sheet = book[name] if book is not None else None
            rows = [
                [str(value) for value in row] for row in frame.itertuples(index=False, name=None)
            ]
            body = _sheet_table(rows, sheet, html_to_markdown)
            facts: dict[str, Any] = (
                _style_facts(sheet, colors)
                if sheet is not None and colors is not None
                else {"profile": STYLE_PROFILE, "style_status": "NOT_INSPECTED"}
            )
            sections.append(
                f"## {markdown_literal(name)}\n{body}\n\n"
                + _style_markdown(facts)
            )
        coverage = "> Excel format coverage: " + STYLE_PROFILE + ". " + " ".join(limitations)
        return coverage + "\n\n" + "\n\n".join(sections).strip()
    finally:
        if book is not None:
            book.close()
